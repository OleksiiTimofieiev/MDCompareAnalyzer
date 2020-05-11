#!/usr/bin/python

from datetime import datetime
import time
# from progress.bar import IncrementalBar
# from progress.bar import Bar
from alive_progress import alive_bar



startTime = datetime.now()

import openpyxl

def convertTuple(tup): 
    str =  ''.join(tup) 
    return str

filename = '4822.xlsx'

## opening the xlsx file
wb = openpyxl.load_workbook('4822.xlsx', read_only=True)
sheet = wb['VerticalDiffs1']
# print(sheet.max_row)

maxval  = sheet.max_row
# maxval = 100000
import sys

# pbar = Bar('Countdown', suffix='%(percent).1f%%', max = maxval)
# pbar.start()


# print('loaded')
## getting the data from the sheet
data = sheet.rows
# print(data)
# print('loaded')
## creating a csv file
# csv = open("data.csv", "w+")
b = 0

z = []

# startTime = datetime.now()

# # https://github.com/rsalmei/alive-progress
with alive_bar(maxval, bar = 'classic', spinner = 'pulse') as bar:
    for row in data:
        l = list(row)
        print(str(l))
        x = ""
        # print(len(l))
        for i in range(len(l)):
            if i == len(l) - 1:
                x += str(l[i].value)
            else:
                x += (str(l[i].value) + ',')

        z.append(x)
        bar()
        # if b == 100000:
            # break
        # b = b + 1
    # # print(x.split(','))
    # # b = b + 1
    # # csv.write('\n')
    # # pbar.next()
# # pbar.finish()
# with alive_bar(maxval, bar = 'classic', spinner = 'pulse') as bar:
    # for row in data:
        # l = list(row)
        # x= ""
        # # print(len(l))
        # for i in range(len(l)):
            # if i == len(l) - 1:
                # x += str(l[i].value)
            # else:
                # x += (str(l[i].value) + ',')

        # z.append(x)
        # bar()

# with alive_bar(maxval, bar = 'classic', spinner = 'pulse') as bar:
    # for row in data:
        # l = list(row)
        # x= ""
        # # print(len(l))
        # for i in range(len(l)):
            # if i == len(l) - 1:
                # x += str(l[i].value)
            # else:
                # x += (str(l[i].value) + ',')

        # z.append(x)
        # bar()

# with alive_bar(maxval, bar = 'classic', spinner = 'pulse') as bar:
    # for row in data:
        # l = list(row)
        # x= ""
        # # print(len(l))
        # for i in range(len(l)):
            # if i == len(l) - 1:
                # x += str(l[i].value)
            # else:
                # x += (str(l[i].value) + ',')

        # z.append(x)
        # bar()

# with alive_bar(maxval, bar = 'classic', spinner = 'pulse') as bar:
    # for row in data:
        # l = list(row)
        # x= ""
        # # print(len(l))
        # for i in range(len(l)):
            # if i == len(l) - 1:
                # x += str(l[i].value)
            # else:
                # x += (str(l[i].value) + ',')

        # z.append(x)
        # bar()

# with alive_bar(maxval, bar = 'classic', spinner = 'pulse') as bar:
    # for row in data:
        # l = list(row)
        # x= ""
        # # print(len(l))
        # for i in range(len(l)):
            # if i == len(l) - 1:
                # x += str(l[i].value)
            # else:
                # x += (str(l[i].value) + ',')

        # z.append(x)
        # bar()        
print(datetime.now() - startTime)

startTime = datetime.now()

for j in z:
    # print(j)
    if "!" in j:
        q = j.split(',')
        print(q)
        if q[1] == 'ACT_FLAG3' and q[4] != '""' and q[5] != 'null':
            print(q)
    
print(datetime.now() - startTime)



# print(x)
# print(z)

## close the csv file
# csv.close()




# with open('data.csv') as file:
    # line = csv.DictReader(file)
    # for row in line:
        # print(row['FID'] ,row ['FID'],row ['FID'])

    
# print(datetime.now() - startTime)

